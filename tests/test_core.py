import os

import pymupdf
import pytest

import core
from core import IslemHatasi


def test_aralik_coz_sirayi_korur_ve_hatayi_bildirir():
    assert core.aralik_coz("4, 1-2", 5) == [3, 0, 1]
    assert core.aralik_coz("2,2,2", 5) == [1]
    for kotu in ("0", "9", "a-b", "3-1", ""):
        with pytest.raises(IslemHatasi):
            core.aralik_coz(kotu, 5)


def test_tc_dogrulama():
    assert core.tc_gecerli("10000000146")
    assert not core.tc_gecerli("12345678901")
    assert not core.tc_gecerli("00000000000")
    assert not core.tc_gecerli("1234")


def test_orijinalin_uzerine_yazilamaz(pdf_yap, ayar):
    yol = pdf_yap(["a"])
    with pytest.raises(IslemHatasi):
        core.sikistir(yol, ayar, yol)


def test_birlestir_yer_imi_ve_icindekiler(pdf_yap, ayar, tmp_path):
    a, b = pdf_yap(["bir", "iki"], "a.pdf"), pdf_yap(["uc"], "b.pdf")
    cikti = str(tmp_path / "m.pdf")
    ayar.update(icindekiler=True, bates=True)
    core.birlestir([a, b], ayar, cikti)
    with pymupdf.open(cikti) as d:
        assert d.page_count == 4                      # 1 içindekiler + 3 sayfa
        assert [t[1] for t in d.get_toc()] == ["İçindekiler", "a", "b"]
        assert len(d[0].get_links()) == 2             # her dosya için tıklanabilir bağlantı
        assert "Sayfa 4 / 4" in d[3].get_text()


def test_tiff_her_kare_ayri_sayfa(tmp_path, ayar):
    from PIL import Image
    kareler = [Image.new("RGB", (300, 400), c) for c in ("red", "green", "blue")]
    tif = str(tmp_path / "s.tif")
    kareler[0].save(tif, save_all=True, append_images=kareler[1:])
    cikti = str(tmp_path / "m.pdf")
    core.birlestir([tif], ayar, cikti)
    with pymupdf.open(cikti) as d:
        assert d.page_count == 3


def test_sayfa_sec_sil_ve_sirala(pdf_yap, tmp_path):
    yol = pdf_yap(["1", "2", "3", "4"])
    cikti = str(tmp_path / "x.pdf")
    assert core.sayfa_sec(yol, "4, 1", cikti) == 2
    with pymupdf.open(cikti) as d:
        assert d[0].get_text().strip() == "4"
    assert core.sayfa_sec(yol, "1-2", cikti, sil=True) == 2
    with pytest.raises(IslemHatasi):
        core.sayfa_sec(yol, "1-4", cikti, sil=True)   # tüm sayfalar silinemez


def test_bos_sayfa_bulma_ve_silme(pdf_yap, tmp_path):
    yol = pdf_yap(["dolu", None, None, "dolu"])
    bos, toplam = core.bos_sayfalari_bul(yol)
    assert (bos, toplam) == ([2, 3], 4)
    cikti = str(tmp_path / "temiz.pdf")
    assert core.sayfalari_sil(yol, bos, cikti) == 2
    assert pymupdf.open(cikti).page_count == 2


@pytest.mark.parametrize("derece", [0, 90, 180, 270])
def test_kase_donmus_sayfada_sayfa_icinde_kalir(pdf_yap, tmp_path, derece):
    yol = pdf_yap(["x"])
    with pymupdf.open(yol) as d:
        d[0].set_rotation(derece)
        d.saveIncr()
    cikti = str(tmp_path / "k.pdf")
    core.kase_bas(yol, cikti, [("ASLI GİBİDİR", None), ("Av. Test", "siyah")], "lacivert", "sag_alt")
    with pymupdf.open(cikti) as d:
        # kaşe metni sayfada okunabilir şekilde bulunmalı
        assert "ASLI" in d[0].get_text() and d[0].rect.width > 0


def test_kase_bos_metin_hatasi(pdf_yap, tmp_path):
    with pytest.raises(IslemHatasi):
        core.kase_bas(pdf_yap(["x"]), str(tmp_path / "k.pdf"), [("", None)], "kirmizi", "sag_alt")


def test_kvkk_tara_karart_ve_dogrula(pdf_yap, tmp_path):
    metin = "TC 10000000146 IBAN TR12 3456 7890 1234 5678 9012 34 tel 0532 123 45 67 a@b.com Gizli"
    yol = pdf_yap([metin, metin])
    rapor = core.kvkk_tara(yol, ["tc", "iban", "tel", "eposta"], ["Gizli"], None, True)
    assert rapor["sayim"] == {"tc": 2, "iban": 2, "tel": 2, "eposta": 2, "ozel": 2}
    cikti = str(tmp_path / "k.pdf")
    adet, metinsiz, kalan = core.karart(yol, cikti, ["tc", "iban", "tel", "eposta"], ["Gizli"], None, True)
    assert adet == 10 and kalan == 0 and metinsiz == []
    with pymupdf.open(cikti) as d:
        yazi = d[0].get_text()
        assert "10000000146" not in yazi and "TR12" not in yazi and "Gizli" not in yazi


def test_kvkk_gecersiz_tc_dogrulama_aciksa_atlanir(pdf_yap):
    yol = pdf_yap(["gecerli 10000000146 gecersiz 12345678901"])
    assert core.kvkk_tara(yol, ["tc"], [], None, True)["sayim"] == {"tc": 1}
    assert core.kvkk_tara(yol, ["tc"], [], None, False)["sayim"] == {"tc": 2}


def test_kvkk_taranmis_sayfa_bildirilir(tmp_path):
    from PIL import Image
    d = pymupdf.open()
    s = d.new_page()
    img = tmp_path / "t.png"
    Image.new("RGB", (200, 200), "white").save(img)
    s.insert_image(pymupdf.Rect(50, 50, 250, 250), filename=str(img))
    yol = str(tmp_path / "scan.pdf")
    d.save(yol)
    assert core.kvkk_tara(yol, ["tc"], [])["metinsiz"] == [1]


def test_ustveri_temizligi(pdf_yap, ayar, tmp_path):
    yol = pdf_yap(["a"])
    with pymupdf.open(yol) as d:
        d.set_metadata({"author": "Gizli Yazar"})
        d.saveIncr()
    cikti = str(tmp_path / "s.pdf")
    core.sikistir(yol, ayar, cikti)
    assert not (pymupdf.open(cikti).metadata or {}).get("author")


def test_boyuta_gore_bol_her_parca_limit_altinda(pdf_yap, tmp_path):
    import random
    import string
    rnd = random.Random(1)                              # sıkışmayan, gerçekçi metin
    yol = pdf_yap(["".join(rnd.choice(string.ascii_letters + " ") for _ in range(2500)) for _ in range(8)])
    klasor = tmp_path / "parcalar"
    klasor.mkdir()
    limit_mb = os.path.getsize(yol) / 1048576 / 3
    sonuc = core.boyuta_gore_bol(yol, limit_mb, str(klasor))
    assert len(sonuc) >= 3
    assert all(mb <= limit_mb for _, mb, _ in sonuc)
    assert sum(pymupdf.open(str(klasor / ad)).page_count for ad, _, _ in sonuc) == 8


def test_turkce_duyarsiz_arama(pdf_yap, tmp_path):
    import docx
    w = docx.Document()
    w.add_paragraph("Vekalet ücreti ŞİKAYET dilekçesi")
    yol = str(tmp_path / "w.docx")
    w.save(yol)
    sonuc, _ = core.icerikte_ara([yol], "sikayet")
    assert len(sonuc) == 1 and "ŞİKAYET" in sonuc[0]["alinti"]
    with pytest.raises(IslemHatasi):
        core.icerikte_ara([yol], "a")


def test_dizin_xlsx(pdf_yap, tmp_path):
    import openpyxl
    yol = pdf_yap(["a"])
    kayit = [{"yol": yol, "tur": "PDF", "sayfa": 1, "mb": 0.01}]
    cikti = str(tmp_path / "d.xlsx")
    assert core.dizin_xlsx(kayit, cikti) == 1
    ws = openpyxl.load_workbook(cikti).active
    assert ws["B2"].value == "a.pdf" and len(ws["G2"].value) == 64


def test_sha256_ayni_icerik(pdf_yap, tmp_path):
    import shutil
    a = pdf_yap(["x"], "a.pdf")
    b = str(tmp_path / "kopya.pdf")
    shutil.copy(a, b)
    assert core.sha256_dosya(a) == core.sha256_dosya(b)


def test_filigran_ve_gorsele_cevir(pdf_yap, tmp_path):
    yol = pdf_yap(["a", "b"])
    assert core.filigran_bas(yol, str(tmp_path / "f.pdf"), "ÖRNEKTİR", renk="mavi") == 2
    klasor = tmp_path / "img"
    klasor.mkdir()
    assert core.gorsele_cevir(yol, str(klasor), "png", 72) == 2
    assert len(os.listdir(klasor)) == 2


def test_word_pdf_windows_disinda_anlasilir_hata(tmp_path):
    import docx
    import sys
    if sys.platform.startswith("win"):
        pytest.skip("Windows'ta Word gerekir")
    w = docx.Document()
    w.add_paragraph("x")
    yol = str(tmp_path / "w.docx")
    w.save(yol)
    with pytest.raises(IslemHatasi):
        core.docx_to_pdf(yol, str(tmp_path / "w.pdf"))
