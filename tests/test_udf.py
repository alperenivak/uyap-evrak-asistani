import docx
import pymupdf
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.shared import Pt

import core
import udf
from conftest import ORNEK_UDF


def test_udf_oku_biçim_ve_hiza(udf_yap):
    bloklar = udf.udf_oku(udf_yap(ORNEK_UDF))
    assert [b["tur"] for b in bloklar] == ["p", "p"]
    assert bloklar[0]["hiza"] == "1" and bloklar[0]["parcalar"][0]["b"] is True
    assert bloklar[1]["parcalar"][0]["i"] is True
    assert udf.paragraf_metni(bloklar[0]) == "Baslik metni"


def test_udf_tablo_okunur(udf_yap):
    xml = ('<?xml version="1.0"?><template><content><![CDATA[A\nB\n]]></content><elements><table><row>'
           '<cell><paragraph><content startOffset="0" length="2"/></paragraph></cell>'
           '<cell><paragraph><content startOffset="2" length="2"/></paragraph></cell></row></table>'
           '</elements></template>')
    b = udf.udf_oku(udf_yap(xml))
    assert b[0]["tur"] == "tablo" and len(b[0]["satirlar"][0]) == 2


def test_bozuk_udf_anlasilir_hata(tmp_path):
    kotu = tmp_path / "x.udf"
    kotu.write_bytes(b"zip degil")
    with pytest.raises(core.IslemHatasi):
        udf.udf_oku(str(kotu))


def test_udf_to_pdf_turkce_karakter(udf_yap, tmp_path):
    xml = ORNEK_UDF.replace("Baslik metni", "ĞÜŞİÖÇ başlık")
    yol = udf_yap(xml.replace('length="13"', 'length="13"'))
    cikti = str(tmp_path / "u.pdf")
    udf.udf_to_pdf(yol, cikti)
    with pymupdf.open(cikti) as d:
        assert d.page_count >= 1 and "ĞÜŞİÖÇ" in d[0].get_text()


def test_udf_to_docx_geri_donusu(udf_yap, tmp_path):
    cikti = str(tmp_path / "u.docx")
    udf.udf_to_docx(udf_yap(ORNEK_UDF), cikti)
    p = docx.Document(cikti).paragraphs
    assert p[0].text == "Baslik metni" and p[0].runs[0].bold


def test_docx_udf_docx_gidis_donus(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.alignment = A.CENTER
    r = p.add_run("T.C. İSTANBUL MAHKEMESİ")
    r.bold, r.font.size = True, Pt(14)
    q = d.add_paragraph("Davacı: Ahmet Şahin ")
    q.add_run("kalın").bold = True
    d.add_paragraph("Emoji 😀 testi")
    tablo = d.add_table(rows=1, cols=2)
    tablo.cell(0, 0).text, tablo.cell(0, 1).text = "Ad", "Değer"
    w = str(tmp_path / "d.docx")
    d.save(w)

    hedef, uyarilar = udf.docx_to_udf(w, str(tmp_path / "d.udf"))
    assert any("tablo" in u for u in uyarilar)
    bloklar = [b for b in udf.udf_oku(hedef) if b["tur"] == "p"]
    metinler = [udf.paragraf_metni(b) for b in bloklar]
    assert metinler[0] == "T.C. İSTANBUL MAHKEMESİ" and bloklar[0]["hiza"] == "1"
    assert metinler[1] == "Davacı: Ahmet Şahin kalın"
    assert "□" in metinler[2]                          # BMP dışı karakter güvenli karakterle değişir
    assert "Ad | Değer" in metinler


def test_birlestirme_udf_dahil(udf_yap, pdf_yap, ayar, tmp_path):
    pdf = pdf_yap(["bir"])
    cikti = str(tmp_path / "m.pdf")
    core.birlestir([pdf, udf_yap(ORNEK_UDF)], ayar, cikti)
    with pymupdf.open(cikti) as d:
        assert d.page_count == 2 and "Baslik" in d[1].get_text()


def test_docx_metni(tmp_path):
    d = docx.Document()
    d.add_paragraph("merhaba")
    yol = str(tmp_path / "w.docx")
    d.save(yol)
    assert udf.docx_metni(yol) == "merhaba"
