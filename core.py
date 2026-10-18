# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Alperen İVAK
"""
UYAP Evrak Asistanı — PDF işlem çekirdeği.
Geliştirici: Alperen İVAK

Bu modül arayüzden bağımsızdır (tkinter içermez); her fonksiyon yolları ve
parametreleri alır, sonucu döndürür. `ilerle(oran, mesaj)` geri çağrısı isteğe
bağlıdır ve arka plan thread'inden çağrılır.
"""
import io
import os
import re
import math
import uuid
import hashlib

import pymupdf
from PIL import Image, ImageOps

# Kalite profilleri: ad -> (dpi, jpeg kalitesi, metin katmanı korunur mu)
PROFILLER = {
    "Akıllı – metni koru (önerilen)": (150, 70, True),
    "Dengeli – görüntüye çevir (150 DPI)": (150, 70, False),
    "Maksimum sıkıştırma (100 DPI)": (100, 50, False),
}
A4 = (595.0, 842.0)
GORSEL_UZANTILARI = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff")


class IslemHatasi(Exception):
    """Kullanıcıya olduğu gibi gösterilebilecek hata."""


# ----------------------------------------------------------------------------
# Yardımcılar
# ----------------------------------------------------------------------------
def _font_bul():
    for yol in ("C:/Windows/Fonts/arial.ttf",
                "/Library/Fonts/Arial.ttf",
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.exists(yol):
            return yol
    return None


FONT_DOSYA = _font_bul()   # Türkçe karakterler (İ, Ş, Ğ) için


def _font_param():
    """Her işlem benzersiz font adı kullanır; böylece daha önce alt kümeye (subset)
    indirilmiş bir fontla çakışıp harflerin kutuya dönmesi önlenir."""
    if FONT_DOSYA:
        return {"fontfile": FONT_DOSYA, "fontname": "tr" + uuid.uuid4().hex[:8]}
    return {"fontname": "helv"}


def _olcu_fontu():
    return pymupdf.Font(fontfile=FONT_DOSYA) if FONT_DOSYA else pymupdf.Font("helv")


def _ilerle(ilerle, oran, mesaj=""):
    if ilerle:
        ilerle(max(0.0, min(1.0, oran)), mesaj)


def pdf_ac(yol):
    try:
        doc = pymupdf.open(yol)
    except Exception as e:
        raise IslemHatasi(f"Dosya açılamadı: {os.path.basename(yol)}\n({e})")
    if doc.needs_pass:
        doc.close()
        raise IslemHatasi(f"Şifreli PDF desteklenmiyor: {os.path.basename(yol)}")
    return doc


def sayfa_sayisi(yol):
    with pdf_ac(yol) as d:
        return len(d)


def hedef_kontrol(kaynak, hedef):
    if os.path.normcase(os.path.abspath(kaynak)) == os.path.normcase(os.path.abspath(hedef)):
        raise IslemHatasi("Orijinal dosyanın üzerine yazılamaz; farklı bir ad seçin.")


def kaydet(doc, hedef, meta_temizle=True):
    if meta_temizle:
        try:
            doc.set_metadata({})
            doc.del_xml_metadata()
        except Exception:
            pass
    try:
        doc.subset_fonts()      # gömülen fontu yalnızca kullanılan harflere indirger
    except Exception:
        pass
    doc.save(hedef, garbage=4, deflate=True, clean=True)


def aralik_coz(metin, toplam):
    """'1-3, 5' -> [0,1,2,4]. Sıra korunur, tekrarlar atılır."""
    sonuc, gorulen = [], set()
    for parca in metin.replace(";", ",").split(","):
        parca = parca.strip()
        if not parca:
            continue
        try:
            if "-" in parca:
                a, b = (int(x) for x in parca.split("-", 1))
            else:
                a = b = int(parca)
        except ValueError:
            raise IslemHatasi(f"Geçersiz sayfa ifadesi: '{parca}'. Örnek: 1-3, 5")
        if a < 1 or b > toplam or a > b:
            raise IslemHatasi(f"'{parca}' geçersiz; sayfa numaraları 1 ile {toplam} arasında olmalı.")
        for i in range(a - 1, b):
            if i not in gorulen:
                gorulen.add(i)
                sonuc.append(i)
    if not sonuc:
        raise IslemHatasi("Sayfa belirtilmedi.")
    return sonuc


# ----------------------------------------------------------------------------
# Birleştirme / sıkıştırma
# ----------------------------------------------------------------------------
def _gorseli_a4e_ekle(doc, yol, kalite):
    """Görseli A4 sayfa(lar)a ortalar. Çok sayfalı TIFF taramalarda her kare ayrı sayfa olur."""
    from PIL import ImageSequence
    w, h = A4
    with Image.open(yol) as kaynak:
        for kare in ImageSequence.Iterator(kaynak):
            img = ImageOps.exif_transpose(kare.copy())          # telefon fotoğraflarının yönü
            if img.mode != "RGB":
                img = img.convert("RGB")
            sayfa = doc.new_page(width=w, height=h)
            marj = 36.0
            oran = min((w - 2 * marj) / img.width, (h - 2 * marj) / img.height)
            cw, ch = img.width * oran, img.height * oran
            x0, y0 = (w - cw) / 2, (h - ch) / 2
            tampon = io.BytesIO()
            img.save(tampon, format="JPEG", quality=kalite, optimize=True)
            sayfa.insert_image(pymupdf.Rect(x0, y0, x0 + cw, y0 + ch), stream=tampon.getvalue())


def _goruntuye_cevir(kaynak, hedef, dpi, kalite, ilerle=None, taban=0.0, pay=1.0):
    n = len(kaynak)
    for i, sayfa in enumerate(kaynak):
        pix = sayfa.get_pixmap(dpi=dpi)
        veri = pix.pil_tobytes(format="JPEG", quality=kalite, optimize=True)
        yeni = hedef.new_page(width=sayfa.rect.width, height=sayfa.rect.height)
        yeni.insert_image(sayfa.rect, stream=veri)
        _ilerle(ilerle, taban + pay * (i + 1) / n, f"Sayfa {i + 1}/{n}")


def _gorselleri_kucult(doc, dpi, kalite):
    """Metin katmanına dokunmadan gömülü görselleri yeniden örnekler."""
    try:
        doc.rewrite_images(dpi_threshold=dpi + 30, dpi_target=dpi, quality=kalite)
    except AttributeError:
        pass   # eski PyMuPDF: yalnızca yapısal sıkıştırma uygulanır


def bates_bas(doc, on_ek):
    toplam = len(doc)
    font = _font_param()
    for i, sayfa in enumerate(doc):
        metin = f"{on_ek} | Sayfa {i + 1} / {toplam}" if on_ek else f"Sayfa {i + 1} / {toplam}"
        w, h = sayfa.rect.width, sayfa.rect.height
        sayfa.insert_textbox(pymupdf.Rect(0, h - 26, w, h - 8), metin, fontsize=8.5,
                             align=pymupdf.TEXT_ALIGN_CENTER, color=(0.2, 0.2, 0.2), **font)


def _icindekiler_ekle(doc, kayitlar):
    """Belgenin başına tıklanabilir içindekiler sayfası ekler.
    kayitlar: [(baslik, baslangic_sayfasi (1 tabanlı), sayfa_adedi)]. Döner: eklenen sayfa sayısı."""
    satir_sayfa = 38
    cp = math.ceil(len(kayitlar) / satir_sayfa)
    for k in range(cp):
        doc.new_page(pno=k, width=A4[0], height=A4[1])
    font, olcu = _font_param(), _olcu_fontu()
    sag, nokta = 545.0, olcu.text_length(".", 10)
    for i, (baslik, bas, adet) in enumerate(kayitlar):
        sayfa, r = doc[i // satir_sayfa], i % satir_sayfa
        if r == 0:
            sayfa.insert_text((50, 60), "İÇİNDEKİLER" + (" (devam)" if i else ""), fontsize=14, **font)
        y = 96 + r * 18
        no = f"{bas + cp}"
        no_x = sag - olcu.text_length(no, 10)
        metin = f"{i + 1}. {baslik}"
        while olcu.text_length(metin, 10) > 360 and len(metin) > 4:
            metin = metin[:-2]
        if metin != f"{i + 1}. {baslik}":
            metin += "…"
        bitis = 50 + olcu.text_length(metin, 10)
        sayfa.insert_text((50, y), metin, fontsize=10, **font)
        sayfa.insert_text((no_x, y), no, fontsize=10, **font)
        nokta_adet = int((no_x - bitis - 12) / nokta)
        if nokta_adet > 0:
            sayfa.insert_text((bitis + 6, y), "." * nokta_adet, fontsize=10, color=(0.55, 0.55, 0.55), **font)
        sayfa.insert_link({"kind": pymupdf.LINK_GOTO, "from": pymupdf.Rect(50, y - 11, sag, y + 4),
                           "page": bas - 1 + cp, "to": pymupdf.Point(0, 40)})
    return cp


def _pdf_kaynak_ekle(yeni, kaynak, metin_koru, dpi, kalite):
    if metin_koru:
        yeni.insert_pdf(kaynak)
    else:
        _goruntuye_cevir(kaynak, yeni, dpi, kalite)


def birlestir(dosyalar, ayar, hedef, ilerle=None):
    """PDF, UDF ve görselleri sırayla tek PDF'te birleştirir. Her dosya için yer imi eklenir;
    ayar["icindekiler"] açıksa başa tıklanabilir içindekiler sayfası konur. UDF belgeleri
    okunabilir kopya olarak PDF'e çevrilir (imzalı aslı değildir)."""
    import udf
    dpi, kalite, metin_koru = PROFILLER[ayar["profil"]]
    yeni = pymupdf.open()
    kayitlar = []
    n = len(dosyalar)
    try:
        for k, yol in enumerate(dosyalar):
            _ilerle(ilerle, k / n * 0.9, f"İşleniyor: {os.path.basename(yol)}")
            baslangic = len(yeni) + 1
            kucuk = yol.lower()
            if kucuk.endswith(".pdf"):
                with pdf_ac(yol) as kaynak:
                    _pdf_kaynak_ekle(yeni, kaynak, metin_koru, dpi, kalite)
            elif kucuk.endswith(".udf"):
                with pymupdf.open(stream=udf.udf_to_pdf_bytes(yol), filetype="pdf") as kaynak:
                    _pdf_kaynak_ekle(yeni, kaynak, metin_koru, dpi, kalite)
            else:
                _gorseli_a4e_ekle(yeni, yol, kalite)
            adet = len(yeni) - baslangic + 1
            if adet > 0:
                kayitlar.append((os.path.splitext(os.path.basename(yol))[0], baslangic, adet))
        if metin_koru:
            _ilerle(ilerle, 0.92, "Görseller optimize ediliyor...")
            _gorselleri_kucult(yeni, dpi, kalite)
        kaydirma = _icindekiler_ekle(yeni, kayitlar) if ayar.get("icindekiler") and kayitlar else 0
        if ayar["bates"]:
            bates_bas(yeni, ayar["on_ek"])
        if kayitlar:
            toc = [[1, "İçindekiler", 1]] if kaydirma else []
            toc += [[1, ad, bas + kaydirma] for ad, bas, _ in kayitlar]
            yeni.set_toc(toc)
        _ilerle(ilerle, 0.97, "Kaydediliyor...")
        kaydet(yeni, hedef, ayar["meta_temizle"])
    finally:
        yeni.close()
    return hedef


def sikistir(yol, ayar, hedef, ilerle=None):
    hedef_kontrol(yol, hedef)
    dpi, kalite, metin_koru = PROFILLER[ayar["profil"]]
    with pdf_ac(yol) as kaynak:
        if metin_koru:
            _ilerle(ilerle, 0.3, "Görseller optimize ediliyor...")
            _gorselleri_kucult(kaynak, dpi, kalite)
            if ayar["bates"]:
                bates_bas(kaynak, ayar["on_ek"])
            _ilerle(ilerle, 0.9, "Kaydediliyor...")
            kaydet(kaynak, hedef, ayar["meta_temizle"])
        else:
            yeni = pymupdf.open()
            try:
                _goruntuye_cevir(kaynak, yeni, dpi, kalite, ilerle, 0.0, 0.9)
                if ayar["bates"]:
                    bates_bas(yeni, ayar["on_ek"])
                kaydet(yeni, hedef, ayar["meta_temizle"])
            finally:
                yeni.close()
    return hedef


# ----------------------------------------------------------------------------
# Bölme / seçme / döndürme / temizleme
# ----------------------------------------------------------------------------
def boyuta_gore_bol(yol, limit_mb, klasor, meta_temizle=True, ilerle=None):
    """Her parça `limit_mb` sınırını aşmayacak şekilde sırayla böler.
    Döner: [(dosya_adi, mb, sinir_asildi_mi)]"""
    limit = limit_mb * 1048576
    kok = os.path.splitext(os.path.basename(yol))[0]
    sonuclar = []
    with pdf_ac(yol) as ana:
        toplam = len(ana)
        ort_sayfa = os.path.getsize(yol) / max(toplam, 1)
        parca, sayac, no = pymupdf.open(), 0, 1

        def parca_yaz():
            nonlocal parca, no
            cikti = os.path.join(klasor, f"{kok}_parca_{no}.pdf")
            kaydet(parca, cikti, meta_temizle)
            parca.close()
            mb = os.path.getsize(cikti) / 1048576
            sonuclar.append((os.path.basename(cikti), mb, mb > limit_mb))
            no += 1

        for i in range(toplam):
            parca.insert_pdf(ana, from_page=i, to_page=i)
            sayac += 1
            _ilerle(ilerle, (i + 1) / toplam * 0.95, f"Sayfa {i + 1}/{toplam}")
            if sayac * ort_sayfa < limit * 0.5:      # tahmin güvenliyse ölçüm yapma (hız)
                continue
            boyut = len(parca.tobytes(garbage=4, deflate=True))
            if boyut > limit and sayac > 1:
                parca.delete_page(sayac - 1)
                parca_yaz()
                parca = pymupdf.open()
                parca.insert_pdf(ana, from_page=i, to_page=i)
                sayac = 1
        if sayac:
            parca_yaz()
        else:
            parca.close()
    return sonuclar


def sayfa_sayfa_bol(yol, klasor, meta_temizle=True, ilerle=None):
    kok = os.path.splitext(os.path.basename(yol))[0]
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i in range(n):
            tek = pymupdf.open()
            try:
                tek.insert_pdf(doc, from_page=i, to_page=i)
                kaydet(tek, os.path.join(klasor, f"{kok}_sayfa_{i + 1}.pdf"), meta_temizle)
            finally:
                tek.close()
            _ilerle(ilerle, (i + 1) / n, f"Sayfa {i + 1}/{n}")
    return n


def sayfa_sec(yol, aralik, hedef, sil=False, meta_temizle=True):
    """sil=False: yalnızca belirtilen sayfaları (yazıldığı sırayla) al.
    sil=True : belirtilen sayfaları çıkar."""
    hedef_kontrol(yol, hedef)
    with pdf_ac(yol) as doc:
        secilen = aralik_coz(aralik, len(doc))
        if sil:
            kalan = [i for i in range(len(doc)) if i not in set(secilen)]
            if not kalan:
                raise IslemHatasi("Tüm sayfalar silinemez.")
            doc.select(kalan)
        else:
            doc.select(secilen)
        kaydet(doc, hedef, meta_temizle)
        return len(doc)


def dondur(yol, aralik, derece, hedef, meta_temizle=True):
    hedef_kontrol(yol, hedef)
    with pdf_ac(yol) as doc:
        idx = aralik_coz(aralik, len(doc)) if aralik.strip() else range(len(doc))
        for i in idx:
            doc[i].set_rotation((doc[i].rotation + derece) % 360)
        kaydet(doc, hedef, meta_temizle)
        return len(list(idx))


def _sayfa_bos_mu(sayfa):
    if sayfa.get_text().strip():
        return False
    pix = sayfa.get_pixmap(dpi=50, colorspace=pymupdf.csGRAY)
    hist = Image.frombytes("L", (pix.width, pix.height), pix.samples).histogram()
    koyu = sum(hist[:215])
    return koyu / (pix.width * pix.height) < 0.002


def bos_sayfalari_bul(yol, ilerle=None):
    """Boş görünen sayfaların numaralarını (1 tabanlı) döndürür; dosyayı değiştirmez."""
    bos = []
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i in range(n):
            if _sayfa_bos_mu(doc[i]):
                bos.append(i + 1)
            _ilerle(ilerle, (i + 1) / n, f"Sayfa {i + 1}/{n} inceleniyor")
    return bos, n


def sayfalari_sil(yol, sayfa_nolari, hedef, meta_temizle=True):
    """sayfa_nolari: 1 tabanlı sayfa numaraları."""
    hedef_kontrol(yol, hedef)
    with pdf_ac(yol) as doc:
        silinecek = {n - 1 for n in sayfa_nolari}
        kalan = [i for i in range(len(doc)) if i not in silinecek]
        if not kalan:
            raise IslemHatasi("Tüm sayfalar silinemez.")
        doc.select(kalan)
        kaydet(doc, hedef, meta_temizle)
    return len(silinecek)


# ----------------------------------------------------------------------------
# Kaşe / filigran
# ----------------------------------------------------------------------------
RENKLER = {"kirmizi": (0.78, 0.05, 0.05), "siyah": (0.05, 0.05, 0.05), "lacivert": (0.0, 0.18, 0.38),
           "mavi": (0.05, 0.25, 0.75), "gri": (0.45, 0.45, 0.45)}
RENK_ADLARI = {"kirmizi": "Kırmızı", "siyah": "Siyah", "lacivert": "Lacivert", "mavi": "Mavi", "gri": "Gri"}


def kase_bas(yol, hedef, satirlar, renk, konum, aralik=None, meta_temizle=True):
    """satirlar: [(metin, renk_adi | None)] — boş metinler atlanır, renk None ise `renk` kullanılır.
    Döndürülmüş sayfalarda da kaşe ekranda görünen sağ/sol/üst/alt köşeye ve dik okunur basılır."""
    hedef_kontrol(yol, hedef)
    satirlar = [(m.strip(), r) for m, r in satirlar if m.strip()]
    if not satirlar:
        raise IslemHatasi("Kaşe için en az bir metin girin.")
    olcu = _olcu_fontu()
    ana_renk = RENKLER.get(renk, RENKLER["kirmizi"])
    punto = 9.0
    genislik = max(150.0, min(300.0, max(olcu.text_length(m, 9) for m, _ in satirlar) + 24))
    while punto > 6 and max(olcu.text_length(m, punto) for m, _ in satirlar) > genislik - 10:
        punto -= 0.5
    satir_h = 13.0
    yukseklik = len(satirlar) * satir_h + 12
    kenar = 25
    font = _font_param()

    with pdf_ac(yol) as doc:
        idx = aralik_coz(aralik, len(doc)) if aralik else list(range(len(doc)))
        for i in idx:
            s = doc[i]
            w, h = s.rect.width, s.rect.height          # ekranda görünen boyut
            x = {"sag_alt": w - genislik - kenar, "sag_ust": w - genislik - kenar,
                 "sol_alt": kenar, "sol_ust": kenar, "alt_orta": (w - genislik) / 2}.get(konum, w - genislik - kenar)
            y = kenar if konum in ("sag_ust", "sol_ust") else h - yukseklik - kenar
            kutu = pymupdf.Rect(x, y, x + genislik, y + yukseklik)
            donus = s.rotation
            dm = s.derotation_matrix                    # görünen -> gerçek sayfa koordinatı
            s.draw_rect(kutu * dm, color=ana_renk, width=1.4)
            for k, (metin, r) in enumerate(satirlar):
                renk_kodu = RENKLER.get(r, ana_renk) if r else ana_renk
                satir = pymupdf.Rect(kutu.x0 + 5, kutu.y0 + 6 + k * satir_h,
                                     kutu.x1 - 5, kutu.y0 + 6 + (k + 1) * satir_h + 3)
                s.insert_textbox(satir * dm, metin, fontsize=punto, color=renk_kodu,
                                 align=pymupdf.TEXT_ALIGN_CENTER, rotate=donus, **font)
        kaydet(doc, hedef, meta_temizle)
        return len(idx)


def filigran_bas(yol, hedef, metin, opaklik=0.25, meta_temizle=True, ilerle=None, renk="gri"):
    hedef_kontrol(yol, hedef)
    metin = metin.strip()
    if not metin:
        raise IslemHatasi("Filigran metni boş olamaz.")
    olcu = _olcu_fontu()
    birim = olcu.text_length(metin, 1)
    font = _font_param()
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i, s in enumerate(doc):
            w, h = s.rect.width, s.rect.height
            punto = max(10.0, min(110.0, math.hypot(w, h) * 0.7 / birim))
            merkez = pymupdf.Point(w / 2, h / 2)
            baslangic = pymupdf.Point(merkez.x - birim * punto / 2, merkez.y + punto / 3)
            s.insert_text(baslangic, metin, fontsize=punto, color=RENKLER.get(renk, RENKLER["gri"]),
                          fill_opacity=opaklik, morph=(merkez, pymupdf.Matrix(45)),
                          overlay=True, **font)
            _ilerle(ilerle, (i + 1) / n * 0.95, f"Sayfa {i + 1}/{n}")
        kaydet(doc, hedef, meta_temizle)
        return n


# ----------------------------------------------------------------------------
# KVKK karartma
# ----------------------------------------------------------------------------
DESENLER = {
    "tc": r"\b[1-9][0-9]{10}\b",
    "iban": r"TR\s*[0-9]{2}(?:\s*[0-9]{4}){5}\s*[0-9]{2}",
    "tel": r"(?:\+90\s*|0\s*)?5[0-9]{2}[\s\-]?[0-9]{3}[\s\-]?[0-9]{2}[\s\-]?[0-9]{2}",
    "eposta": r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}",
}


ETIKETLER = {"tc": "T.C. kimlik no", "iban": "IBAN", "tel": "Telefon", "eposta": "E-posta",
             "ozel": "Özel kelime"}


def tc_gecerli(no):
    """11 haneli T.C. kimlik numarası algoritma doğrulaması."""
    if len(no) != 11 or not no.isdigit() or no[0] == "0":
        return False
    r = [int(c) for c in no]
    if ((sum(r[0:9:2]) * 7) - sum(r[1:8:2])) % 10 != r[9]:
        return False
    return sum(r[:10]) % 10 == r[10]


def _eslesmeler(metin, desenler, ozel, tc_dogrula=False):
    """Sayfa metnindeki eşleşmeleri türe göre listeler: {tur: [eşleşme, ...]}"""
    sonuc = {}
    for d in desenler:
        liste = [" ".join(e.group().split()) for e in re.finditer(DESENLER[d], metin)]
        if d == "tc" and tc_dogrula:
            liste = [x for x in liste if tc_gecerli(x)]
        if liste:
            sonuc[d] = liste
    kucuk = metin.lower()
    liste = []
    for k in ozel:
        liste += [k] * kucuk.count(k.lower())
    if liste:
        sonuc["ozel"] = liste
    return sonuc


def _metinsiz_mi(sayfa):
    return not sayfa.get_text().strip() and bool(sayfa.get_images())


def kvkk_tara(yol, desenler, ozel, ilerle=None, tc_dogrula=False):
    """Dosyayı DEĞİŞTİRMEDEN kişisel veri eşleşmelerini raporlar."""
    sayim, sayfalar, metinsiz = {}, {}, []
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i, sayfa in enumerate(doc):
            metin = sayfa.get_text()
            if not metin.strip():
                if sayfa.get_images():
                    metinsiz.append(i + 1)
            else:
                for tur, liste in _eslesmeler(metin, desenler, ozel, tc_dogrula).items():
                    sayim[tur] = sayim.get(tur, 0) + len(liste)
                    sayfalar.setdefault(tur, []).append(i + 1)
            _ilerle(ilerle, (i + 1) / n, f"Sayfa {i + 1}/{n} taranıyor")
    return {"sayim": sayim, "sayfalar": sayfalar, "metinsiz": metinsiz}


def karart(yol, hedef, desenler, ozel_kelimeler, ilerle=None, tc_dogrula=False):
    """Seçili veri türlerini ve özel kelimeleri kalıcı olarak siler, ardından çıktıyı
    yeniden tarayarak doğrular.
    Döner: (maskelenen_adet, metin_katmani_olmayan_sayfalar, çıktıda_kalan_eşleşme)"""
    hedef_kontrol(yol, hedef)
    adet, metinsiz = 0, []
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i, sayfa in enumerate(doc):
            metin = sayfa.get_text()
            if not metin.strip():
                if sayfa.get_images():
                    metinsiz.append(i + 1)       # taranmış sayfa: metin bulunamaz
                continue
            aranacaklar = set()
            for liste in _eslesmeler(metin, desenler, ozel_kelimeler, tc_dogrula).values():
                aranacaklar.update(liste)
            kutular = set()
            for a in aranacaklar:
                for r in sayfa.search_for(a):
                    kutular.add((r.x0, r.y0, r.x1, r.y1))
            for k in kutular:
                sayfa.add_redact_annot(pymupdf.Rect(*k), fill=(0, 0, 0))
            if kutular:
                sayfa.apply_redactions()
                adet += len(kutular)
            _ilerle(ilerle, (i + 1) / n * 0.9, f"Sayfa {i + 1}/{n}")
        try:
            doc.scrub(remove_links=False)     # gizli metin, meta veri, eklenti vb.
        except Exception:
            pass
        kaydet(doc, hedef, True)
    _ilerle(ilerle, 0.95, "Çıktı doğrulanıyor...")
    kalan = 0
    with pdf_ac(hedef) as cikti:
        for sayfa in cikti:
            kalan += sum(len(v) for v in _eslesmeler(sayfa.get_text(), desenler, ozel_kelimeler, tc_dogrula).values())
    return adet, metinsiz, kalan


# ----------------------------------------------------------------------------
# Dönüştürme / metin
# ----------------------------------------------------------------------------
def gorsele_cevir(yol, klasor, bicim="png", dpi=150, ilerle=None):
    kok = os.path.splitext(os.path.basename(yol))[0]
    with pdf_ac(yol) as doc:
        n = len(doc)
        for i, s in enumerate(doc):
            pix = s.get_pixmap(dpi=dpi)
            cikti = os.path.join(klasor, f"{kok}_sayfa_{i + 1}.{bicim}")
            pix.save(cikti, jpg_quality=88) if bicim == "jpg" else pix.save(cikti)
            _ilerle(ilerle, (i + 1) / n, f"Sayfa {i + 1}/{n}")
    return n


def metin_cikar(yol):
    parcalar = []
    with pdf_ac(yol) as doc:
        for i, s in enumerate(doc):
            t = s.get_text().strip()
            if t:
                parcalar.append(f"--- Sayfa {i + 1} ---\n{t}")
    return "\n\n".join(parcalar)


def docx_to_pdf(yol, hedef):
    """Word belgesini PDF'e çevirir. Windows'ta Microsoft Word kurulu olmalıdır (docx2pdf)."""
    import sys
    hedef_kontrol(yol, hedef)
    try:
        import docx2pdf
    except ImportError:
        raise IslemHatasi("Word ➔ PDF için 'docx2pdf' paketi kurulu değil.")
    for ad in ("stdout", "stderr"):                  # pencereli (--noconsole) EXE'de akışlar None olur
        if getattr(sys, ad) is None:
            setattr(sys, ad, open(os.devnull, "w"))
    com = None
    try:
        import pythoncom
        pythoncom.CoInitialize()                     # thread içinde Word'e erişim için
        com = pythoncom
    except ImportError:
        pass
    try:
        docx2pdf.convert(yol, hedef)
    except NotImplementedError:
        raise IslemHatasi("Word ➔ PDF yalnızca Windows ve macOS'ta, Microsoft Word kuruluyken çalışır.")
    except Exception as e:
        raise IslemHatasi("Word ➔ PDF dönüştürmesi başarısız oldu. Microsoft Word'ün kurulu ve lisanslı "
                          f"olduğundan, belgenin başka bir programda açık olmadığından emin olun.\n({e})")
    finally:
        if com:
            com.CoUninitialize()
    if not os.path.isfile(hedef):
        raise IslemHatasi("Dönüştürme tamamlandı ancak PDF dosyası oluşturulamadı.")
    return hedef


# ----------------------------------------------------------------------------
# Aynı içerik tespiti, içerikte arama, Excel evrak dizini
# ----------------------------------------------------------------------------
def sha256_dosya(yol):
    h = hashlib.sha256()
    with open(yol, "rb") as f:
        for blok in iter(lambda: f.read(1 << 20), b""):
            h.update(blok)
    return h.hexdigest()


_TR = str.maketrans({"İ": "i", "I": "i", "ı": "i", "Ş": "s", "ş": "s", "Ğ": "g", "ğ": "g",
                     "Ü": "u", "ü": "u", "Ö": "o", "ö": "o", "Ç": "c", "ç": "c"})


def tr_sadelestir(metin):
    """Büyük/küçük harf ve Türkçe karakter farkını yok sayan, uzunluğu korunan sadeleştirme."""
    return metin.translate(_TR).lower()


def _arama_kaynaklari(yol):
    """(yer_adi, metin) çiftlerini üretir."""
    import udf
    kucuk = yol.lower()
    if kucuk.endswith(".pdf"):
        with pdf_ac(yol) as doc:
            for i, s in enumerate(doc):
                yield f"Sayfa {i + 1}", s.get_text(), (not s.get_text().strip() and bool(s.get_images()))
    elif kucuk.endswith(".udf"):
        yield "UDF", udf.udf_metni(yol), False
    elif kucuk.endswith(".docx"):
        yield "Word", udf.docx_metni(yol), False


def icerikte_ara(dosyalar, sorgu, ilerle=None, en_fazla=500):
    """Listedeki PDF/UDF/Word dosyalarının içinde arar; Türkçe karakter ve büyük/küçük harf fark etmez.
    Döner: (sonuçlar[{dosya, yer, alinti}], metni_olmayan_dosyalar)"""
    q = tr_sadelestir(sorgu.strip())
    if len(q) < 2:
        raise IslemHatasi("Aranacak ifade en az 2 karakter olmalı.")
    sonuc, taranmis = [], []
    n = len(dosyalar)
    for k, yol in enumerate(dosyalar):
        _ilerle(ilerle, k / n, f"Aranıyor: {os.path.basename(yol)}")
        metinsiz_sayfa = 0
        try:
            for yer, metin, metinsiz in _arama_kaynaklari(yol):
                metinsiz_sayfa += bool(metinsiz)
                t = tr_sadelestir(metin)
                if len(t) != len(metin):
                    t, metin = t, t
                poz = t.find(q)
                while poz != -1 and len(sonuc) < en_fazla:
                    alinti = " ".join(metin[max(0, poz - 45): poz + len(q) + 45].split())
                    sonuc.append({"dosya": os.path.basename(yol), "yer": yer, "alinti": alinti})
                    poz = t.find(q, poz + len(q))
        except IslemHatasi:
            continue
        if metinsiz_sayfa:
            taranmis.append(os.path.basename(yol))
        if len(sonuc) >= en_fazla:
            break
    return sonuc, taranmis


def dizin_xlsx(kayitlar, hedef, ilerle=None):
    """kayitlar: [{yol, tur, sayfa, mb, hash?}]. Sıra, ad, tür, sayfa, boyut, tarih, SHA-256 içeren
    filtrelenebilir bir Excel evrak dizini yazar; dosya adları belgeyi açan bağlantıdır."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise IslemHatasi("Excel dizini için 'openpyxl' paketi kurulu değil.")
    import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "Evrak Dizini"
    ws.append(["Sıra", "Dosya adı", "Tür", "Sayfa", "Boyut (MB)", "Son değişiklik", "SHA-256"])
    for hucre in ws[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="0B1F3A")
        hucre.alignment = Alignment(horizontal="center")
    n = len(kayitlar)
    for i, k in enumerate(kayitlar, 1):
        _ilerle(ilerle, i / n, f"Dizin: {os.path.basename(k['yol'])}")
        tarih = datetime.datetime.fromtimestamp(os.path.getmtime(k["yol"])).strftime("%d.%m.%Y %H:%M")
        ws.append([i, os.path.basename(k["yol"]), k["tur"], k["sayfa"], round(k["mb"], 2), tarih,
                   k.get("hash") or sha256_dosya(k["yol"])])
        c = ws.cell(row=i + 1, column=2)
        c.hyperlink = os.path.abspath(k["yol"])
        c.font = Font(color="0563C1", underline="single")
    for sut, genislik in zip("ABCDEFG", (6, 56, 8, 8, 12, 18, 66)):
        ws.column_dimensions[sut].width = genislik
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(hedef)
    return n
