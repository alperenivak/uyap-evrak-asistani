# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Alperen İVAK
"""
UYAP Evrak Asistanı — UDF / Word belge dönüştürücüleri.
Geliştirici: Alperen İVAK

UDF, içinde `content.xml` bulunan bir ZIP paketidir. Metin <content> düğümünde tek parça
durur; <elements> altındaki paragraf/tablo düğümleri bu metne `startOffset` ve `length`
ile işaret eder. `udf_oku` bu yapıyı ortak bir blok modeline çevirir; Word, PDF ve ekran
görüntüleyici bu modeli kullanır.

Blok modeli:
  {"tur": "p", "hiza": "0|1|2|3", "parcalar": [{"m", "b", "i", "u", "boyut", "aile", "resim"}]}
  {"tur": "tablo", "satirlar": [[[paragraf, ...], ...], ...]}      # satır > hücre > paragraflar
  {"tur": "sayfa_sonu"}
"""
import base64
import html
import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr

import pymupdf

from core import IslemHatasi

SIFIR_GENISLIK = "\u200b"
HIZA_CSS = {"0": "left", "1": "center", "2": "right", "3": "justify"}


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------
def _xml_oku(yol):
    try:
        with zipfile.ZipFile(yol, "r") as zf:
            if "content.xml" not in zf.namelist():
                raise IslemHatasi("Geçersiz UDF: 'content.xml' bulunamadı.")
            return ET.fromstring(zf.read("content.xml"))
    except zipfile.BadZipFile:
        raise IslemHatasi("UDF dosyası okunamadı (bozuk veya desteklenmeyen biçim).")
    except ET.ParseError:
        raise IslemHatasi("UDF içeriği ayrıştırılamadı.")


def _ana_metin(kok):
    dugum = kok.find("content")
    return (dugum.text or "") if dugum is not None else ""


def _p_oku(el, metin):
    parcalar = []
    for c in el:
        if c.tag == "content":
            bas, uz = int(c.get("startOffset", 0)), int(c.get("length", 0))
            m = metin[bas:bas + uz].replace(SIFIR_GENISLIK, "").replace("\n", "")
            if not m:
                continue
            boyut = None
            try:
                boyut = float(c.get("size")) if c.get("size") else None
            except ValueError:
                pass
            parcalar.append({"m": m, "b": c.get("bold") == "true", "i": c.get("italic") == "true",
                             "u": c.get("underline") == "true", "boyut": boyut, "aile": c.get("family")})
        elif c.tag == "tab":
            parcalar.append({"m": "\t"})
        elif c.tag == "space":
            parcalar.append({"m": " "})
        elif c.tag == "image":
            veri = c.get("imageData") or c.get("data")
            if veri:
                try:
                    parcalar.append({"m": "", "resim": base64.b64decode(veri)})
                except Exception:
                    pass
    return {"tur": "p", "hiza": str(el.get("Alignment", el.get("alignment", "0"))), "parcalar": parcalar}


def _tablo_oku(el, metin):
    satirlar = []
    for satir in el.findall("row"):
        satirlar.append([[_p_oku(p, metin) for p in h.findall("paragraph")] for h in satir.findall("cell")])
    return {"tur": "tablo", "satirlar": satirlar}


def udf_metni(yol):
    """UDF içindeki düz metni döndürür."""
    kok = _xml_oku(yol)
    metin = _ana_metin(kok)
    if not metin.strip():
        metin = "".join(kok.itertext())
    return metin.replace(SIFIR_GENISLIK, "")


def udf_oku(yol):
    """UDF'yi blok modeline çevirir. Biçim bilgisi okunamazsa düz metin satırlarına döner."""
    kok = _xml_oku(yol)
    metin = _ana_metin(kok)
    bloklar = []
    el = kok.find("elements")
    if metin and el is not None:
        for e in el:
            if e.tag == "paragraph":
                bloklar.append(_p_oku(e, metin))
            elif e.tag == "table":
                bloklar.append(_tablo_oku(e, metin))
            elif e.tag in ("pageBreak", "page-break"):
                bloklar.append({"tur": "sayfa_sonu"})
    if not any(b["tur"] in ("p", "tablo") for b in bloklar):
        bloklar = [{"tur": "p", "hiza": "0", "parcalar": [{"m": s}]} for s in udf_metni(yol).splitlines()]
    return bloklar


def paragraf_metni(p):
    return "".join(x["m"] for x in p["parcalar"])


# ----------------------------------------------------------------------------
# UDF -> Word
# ----------------------------------------------------------------------------
def _docx_paragraf_yaz(par, p):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.shared import Pt, Inches
    par.alignment = {"0": A.LEFT, "1": A.CENTER, "2": A.RIGHT, "3": A.JUSTIFY}.get(p["hiza"], A.LEFT)
    for x in p["parcalar"]:
        if x.get("resim"):
            try:
                par.add_run().add_picture(io.BytesIO(x["resim"]), width=Inches(2.5))
            except Exception:
                pass
            continue
        r = par.add_run(x["m"])
        r.bold, r.italic, r.underline = x.get("b") or None, x.get("i") or None, x.get("u") or None
        if x.get("boyut"):
            r.font.size = Pt(x["boyut"])
        if x.get("aile"):
            r.font.name = x["aile"]


def udf_to_docx(yol, hedef):
    """UDF'yi Word'e çevirir (paragraf, hizalama, kalın/italik/altı çizili, tablo, görsel)."""
    import docx
    belge = docx.Document()
    for b in udf_oku(yol):
        if b["tur"] == "p":
            _docx_paragraf_yaz(belge.add_paragraph(), b)
        elif b["tur"] == "tablo" and b["satirlar"]:
            sutun = max(len(s) for s in b["satirlar"]) or 1
            t = belge.add_table(rows=len(b["satirlar"]), cols=sutun)
            t.style = "Table Grid"
            for r, satir in enumerate(b["satirlar"]):
                for c, paragraflar in enumerate(satir):
                    hucre = t.cell(r, c)
                    for k, p in enumerate(paragraflar):
                        _docx_paragraf_yaz(hucre.paragraphs[0] if k == 0 else hucre.add_paragraph(), p)
        elif b["tur"] == "sayfa_sonu":
            belge.add_page_break()
    belge.save(hedef)
    return hedef


# ----------------------------------------------------------------------------
# UDF -> PDF (biçimli, Türkçe karakter destekli)
# ----------------------------------------------------------------------------
def _font_ailesi():
    """(klasör, {'r','b','i','bi'} dosya adları). Önce Times New Roman, sonra Arial, sonra DejaVu."""
    adaylar = [
        ("C:/Windows/Fonts", {"r": "times.ttf", "b": "timesbd.ttf", "i": "timesi.ttf", "bi": "timesbi.ttf"}),
        ("C:/Windows/Fonts", {"r": "arial.ttf", "b": "arialbd.ttf", "i": "ariali.ttf", "bi": "arialbi.ttf"}),
        ("/usr/share/fonts/truetype/dejavu", {"r": "DejaVuSerif.ttf", "b": "DejaVuSerif-Bold.ttf",
                                              "i": "DejaVuSerif-Italic.ttf", "bi": "DejaVuSerif-BoldItalic.ttf"}),
        ("/usr/share/fonts/truetype/dejavu", {"r": "DejaVuSans.ttf", "b": "DejaVuSans-Bold.ttf",
                                              "i": "DejaVuSans-Oblique.ttf", "bi": "DejaVuSans-BoldOblique.ttf"}),
    ]
    for klasor, d in adaylar:
        if os.path.exists(os.path.join(klasor, d["r"])):
            return klasor, {k: v for k, v in d.items() if os.path.exists(os.path.join(klasor, v))}
    return None, {}


def _span_html(x, arc, sayac):
    if x.get("resim"):
        try:
            from PIL import Image
            with Image.open(io.BytesIO(x["resim"])) as im:
                tampon = io.BytesIO()
                im.convert("RGB").save(tampon, format="PNG")
            sayac[0] += 1
            ad = f"resim{sayac[0]}.png"
            arc.add(tampon.getvalue(), ad)
            return f'<img src="{ad}" style="width:200pt">'
        except Exception:
            return ""
    m = html.escape(x["m"])
    m = re.sub(r" {2,}", lambda e: "&nbsp;" * len(e.group()), m).replace("\t", "&nbsp;" * 8)
    stil = []
    if x.get("b"):
        stil.append("font-weight:bold")
    if x.get("i"):
        stil.append("font-style:italic")
    if x.get("u"):
        stil.append("text-decoration:underline")
    if x.get("boyut"):
        stil.append(f"font-size:{x['boyut']}pt")
    return f'<span style="{";".join(stil)}">{m}</span>' if stil else m


def _p_html(p, arc, sayac):
    ic = "".join(_span_html(x, arc, sayac) for x in p["parcalar"]) or "&nbsp;"
    return f'<p style="text-align:{HIZA_CSS.get(p["hiza"], "left")};margin:0 0 3pt 0">{ic}</p>'


def udf_to_pdf_bytes(yol):
    """UDF'nin PDF olarak okunabilir kopyasını üretir. İmzalı aslı DEĞİLDİR."""
    klasor, fontlar = _font_ailesi()
    arc = pymupdf.Archive()
    css = "body{font-size:12pt;}"
    if klasor:
        arc.add(klasor)
        faces = {"r": "", "b": "font-weight:bold;", "i": "font-style:italic;", "bi": "font-weight:bold;font-style:italic;"}
        css = "".join(f"@font-face{{font-family:trf;{faces[k]}src:url({v});}}" for k, v in fontlar.items())
        css += "body{font-family:trf;font-size:12pt;}"
    sayac = [0]
    govde = []
    for b in udf_oku(yol):
        if b["tur"] == "p":
            govde.append(_p_html(b, arc, sayac))
        elif b["tur"] == "tablo":
            satirlar = "".join(
                "<tr>" + "".join("<td>" + "".join(_p_html(p, arc, sayac) for p in h) + "</td>" for h in s) + "</tr>"
                for s in b["satirlar"])
            govde.append(f'<table border="1" cellpadding="3" style="width:100%;border-collapse:collapse">{satirlar}</table>')
        elif b["tur"] == "sayfa_sonu":
            govde.append('<div style="page-break-after:always"></div>')
    story = pymupdf.Story(html="<body>" + "".join(govde) + "</body>", user_css=css, archive=arc)
    tampon = io.BytesIO()
    yazici = pymupdf.DocumentWriter(tampon)
    sayfa = pymupdf.paper_rect("a4")
    alan = sayfa + (42.5, 56, -42.5, -56)
    devam = 1
    while devam:
        dev = yazici.begin_page(sayfa)
        devam, _ = story.place(alan)
        story.draw(dev)
        yazici.end_page()
    yazici.close()
    return tampon.getvalue()


def udf_to_pdf(yol, hedef):
    with open(hedef, "wb") as f:
        f.write(udf_to_pdf_bytes(yol))
    return hedef


# ----------------------------------------------------------------------------
# Word -> UDF (DENEYSEL)
# ----------------------------------------------------------------------------
def _bmp_temiz(s):
    return "".join(c if ord(c) <= 0xFFFF else "□" for c in s)


def docx_to_udf(yol, hedef):
    """Word'ü UDF'ye çevirir. Metin, paragraf hizası, kalın/italik/altı çizili, punto, yazı tipi ve
    sayfa kenar boşlukları taşınır; tablolar satır satır metne, görseller atlanarak çevrilir.
    Dönüş: (hedef, uyarılar). Çıktı yazıldıktan sonra yeniden okunup doğrulanır."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        belge = docx.Document(yol)
    except Exception as e:
        raise IslemHatasi(f"Word belgesi açılamadı: {e}")

    uyarilar, parcalar_metin, elementler, beklenen = [], [], [], []
    ofset = 0
    hiza_kodu = {A.LEFT: "0", A.CENTER: "1", A.RIGHT: "2", A.JUSTIFY: "3"}

    def paragraf(par, duz=None):
        nonlocal ofset
        if duz is not None:
            hiza = None
        else:
            hiza = par.alignment if par.alignment is not None else par.style.paragraph_format.alignment
        girdiler, metin_tum = [], ""
        calismalar = [(duz, None)] if duz is not None else [(r.text, r) for r in par.runs]
        for t, r in calismalar:
            t = _bmp_temiz((t or "").replace("\t", "    ").replace("\n", " "))
            if not t:
                continue
            nitelik = {"startOffset": ofset, "length": len(t)}
            if r is not None:
                f = r.font
                boyut = f.size.pt if f.size else (par.style.font.size.pt if par.style.font.size else 12)
                nitelik.update(size=boyut)
                nitelik.update(family=f.name or par.style.font.name or "Times New Roman")
                if r.bold or (r.bold is None and par.style.font.bold):
                    nitelik["bold"] = "true"
                if r.italic or (r.italic is None and par.style.font.italic):
                    nitelik["italic"] = "true"
                if r.underline:
                    nitelik["underline"] = "true"
            else:
                nitelik.update(size=12, family="Times New Roman")
            girdiler.append(nitelik)
            parcalar_metin.append(t)
            metin_tum += t
            ofset += len(t)
        if not girdiler:
            girdiler.append({"startOffset": ofset, "length": 0, "size": 12, "family": "Times New Roman"})
        girdiler[-1]["length"] += 1            # paragraf sonu satır başı karakteri son parçaya dahildir
        parcalar_metin.append("\n")
        ofset += 1
        beklenen.append(metin_tum)
        ic = "".join("<content " + " ".join(f"{k}={quoteattr(str(v))}" for k, v in g.items()) + " />"
                     for g in girdiler)
        ek = ""
        pf = par.paragraph_format if duz is None else None
        if pf is not None:
            if pf.left_indent:
                ek += f' LeftIndent="{pf.left_indent.pt:.1f}"'
            if pf.first_line_indent:
                ek += f' FirstLineIndent="{pf.first_line_indent.pt:.1f}"'
        elementler.append(f'<paragraph Alignment="{hiza_kodu.get(hiza, "0")}"{ek}>{ic}</paragraph>')

    tablo_var = gorsel = 0
    for c in belge.element.body.iterchildren():
        if c.tag == qn("w:p"):
            par = Paragraph(c, belge)
            if c.findall(".//" + qn("w:drawing")) or c.findall(".//" + qn("w:pict")):
                gorsel += 1
            paragraf(par)
        elif c.tag == qn("w:tbl"):
            tablo_var += 1
            for satir in Table(c, belge).rows:
                gorulen, hucreler = set(), []
                for h in satir.cells:
                    if id(h._tc) not in gorulen:
                        gorulen.add(id(h._tc))
                        hucreler.append(h.text.replace("\n", " ").strip())
                paragraf(None, duz=" | ".join(hucreler))
    if tablo_var:
        uyarilar.append(f"{tablo_var} tablo, hücreler ' | ' ile ayrılmış metin satırlarına çevrildi.")
    if gorsel:
        uyarilar.append(f"{gorsel} görsel taşınmadı.")
    if any(p.text.strip() for s in belge.sections for p in list(s.header.paragraphs) + list(s.footer.paragraphs)):
        uyarilar.append("Üstbilgi/altbilgi taşınmadı.")

    s0 = belge.sections[0]
    sayfa = (f'<pageFormat mediaSizeName="1" leftMargin="{s0.left_margin.pt:.3f}" '
             f'rightMargin="{s0.right_margin.pt:.3f}" topMargin="{s0.top_margin.pt:.3f}" '
             f'bottomMargin="{s0.bottom_margin.pt:.3f}" paperOrientation="1" headerFOffset="20.0" '
             'footerFOffset="20.0" />') if s0.left_margin is not None else ""
    metin = "".join(parcalar_metin).replace("]]>", "]] >")
    xml = ('<?xml version="1.0" encoding="UTF-8" ?>\n<template format_id="1.8" >\n'
           f'<content><![CDATA[{metin}]]></content>\n<properties>{sayfa}</properties>\n'
           f'<elements resolver="hvl-default">\n' + "\n".join(elementler) + '\n</elements>\n'
           '<styles><style name="default" description="Geçerli" family="Times New Roman" size="12" '
           'bold="false" italic="false" foreground="-13421773" />'
           '<style name="hvl-default" family="Times New Roman" size="12" description="Gövde" /></styles>\n'
           '</template>')
    with zipfile.ZipFile(hedef, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.xml", xml.encode("utf-8"))

    # Doğrulama: yazılan UDF yeniden okunur ve paragraf metinleri kaynakla karşılaştırılır
    try:
        okunan = [paragraf_metni(b) for b in udf_oku(hedef) if b["tur"] == "p"]
    except IslemHatasi:
        okunan = None
    if okunan != [b.replace(SIFIR_GENISLIK, "") for b in beklenen]:
        try:
            os.remove(hedef)
        except OSError:
            pass
        raise IslemHatasi("Doğrulama başarısız: üretilen UDF kaynakla uyuşmadı, dosya yazılmadı.")
    return hedef, uyarilar


# ----------------------------------------------------------------------------
# Word metni
# ----------------------------------------------------------------------------
def docx_metni(yol):
    """Word belgesindeki paragraf ve tablo metnini döndürür."""
    import docx
    try:
        belge = docx.Document(yol)
    except Exception as e:
        raise IslemHatasi(f"Word belgesi açılamadı: {e}")
    parcalar = [p.text for p in belge.paragraphs]
    for t in belge.tables:
        for satir in t.rows:
            parcalar.append(" | ".join(h.text.strip() for h in satir.cells))
    return "\n".join(parcalar)
